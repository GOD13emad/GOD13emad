#!/usr/bin/env python3
"""Fictional CSV/SQL-to-Excel ONE-report pilot. No customer data or SQL connection.

Requires openpyxl (pip install openpyxl). Formulas calculate when opened in Excel
or a compatible recalculation engine; independent Python checks run immediately.
"""
from __future__ import annotations
import argparse
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.properties import CalcProperties

ROWS = [
    ("2026-10-01", "Parts", "SYN-001", 2, 480000),
    ("2026-10-02", "Services", "SYN-002", 1, 670000),
    ("2026-10-03", "Tools", "SYN-003", 3, 120000),
    ("2026-10-04", "Parts", "SYN-004", 1, 950000),
    ("2026-10-06", "Services", "SYN-005", 2, 375000),
    ("2026-10-07", "Tools", "SYN-006", 4, 98000),
    ("2026-10-09", "Parts", "SYN-007", 2, 320000),
    ("2026-10-11", "Services", "SYN-008", 3, 225000),
    ("2026-10-12", "Tools", "SYN-009", 1, 710000),
    ("2026-10-14", "Parts", "SYN-010", 5, 89000),
    ("2026-10-18", "Services", "SYN-011", 1, 1250000),
    ("2026-10-21", "Tools", "SYN-012", 2, 145000),
]
EXPECTED_TOMAN = 8092000
EXPECTED_UNITS = 27

def verify_inputs():
    assert len(ROWS) == 12
    assert len(set(r[2] for r in ROWS)) == len(ROWS), "Duplicate synthetic references"
    assert sum(r[3] for r in ROWS) == EXPECTED_UNITS
    assert sum(r[3] * r[4] for r in ROWS) == EXPECTED_TOMAN

def build_workbook(output: Path):
    verify_inputs()
    wb = Workbook()
    guide = wb.active
    guide.title = "Scope_And_Readme"
    source = wb.create_sheet("Synthetic_Inputs")
    report = wb.create_sheet("One_Report_Pilot")
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)

    guide.append(["Synthetic Python-to-Excel reporting proof — NOT customer data"])
    for item in [
        ("Input", "12 fictional transaction records; no customer's database"),
        ("Output", "One category-level report and three reconciliation formulas"),
        ("Python validation", "Independent gross, quantity and unique reference checks"),
        ("Unimplemented", "SQL Server connection, scheduling, four customer report templates"),
        ("Contract gate", "Client-authorized inputs, written acceptance and funded milestone"),
        ("Expected gross", EXPECTED_TOMAN),
        ("Expected units", EXPECTED_UNITS),
    ]: guide.append(item)

    source.append(["Only synthetic records: never use as evidence of earned revenue"])
    source.append([])
    source.append(["Date", "Category", "Reference", "Quantity", "Unit Price (toman)", "Gross (toman)", "State"])
    for idx, (date, category, ref, qty, price) in enumerate(ROWS, start=4):
        source.append([date, category, ref, qty, price, f"=D{idx}*E{idx}", "Posted"])
    source.cell(17, 1, "Synthetic gross total")
    source.cell(17, 6, "=SUM(F4:F15)")

    report.append(["One-report synthetic example — not a client delivery"])
    report.append([])
    report.append(["Category", "Record count", "Unit count", "Gross sales (toman)"])
    for idx, category in enumerate(("Parts", "Services", "Tools"), start=4):
        report.append([
            category,
            f"=COUNTIF(Synthetic_Inputs!$B$4:$B$15,A{idx})",
            f"=SUMIF(Synthetic_Inputs!$B$4:$B$15,A{idx},Synthetic_Inputs!$D$4:$D$15)",
            f"=SUMIF(Synthetic_Inputs!$B$4:$B$15,A{idx},Synthetic_Inputs!$F$4:$F$15)",
        ])
    report.cell(8, 1, "All categories")
    for col, letter in ((2, "B"), (3, "C"), (4, "D")):
        report.cell(8, col, f"=SUM({letter}4:{letter}6)")
    report.cell(11, 1, "Gross total reconciles")
    report.cell(11, 2, '=IF(Synthetic_Inputs!F17=D8,"PASS","FAIL")')
    report.cell(12, 1, "Row count reconciles")
    report.cell(12, 2, '=IF(COUNTA(Synthetic_Inputs!C4:C15)=B8,"PASS","FAIL")')
    report.cell(13, 1, "References unique")
    report.cell(13, 2, '=IF(SUMPRODUCT(1/COUNTIF(Synthetic_Inputs!C4:C15,Synthetic_Inputs!C4:C15))=12,"PASS","FAIL")')
    report.cell(15, 1, "Formulas are recalculated when opened in a spreadsheet application")

    navy, pale = "17456C", "EDF3F8"
    for ws in (guide, source, report):
        ws.freeze_panes = "A4"
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 28
        ws["A1"].font = Font(name="Aptos", size=14, bold=True, color=navy)
        ws.column_dimensions["A"].width = 31
    for ws, row, count in ((source, 3, 7), (report, 3, 4)):
        for col in range(1, count + 1):
            cell = ws.cell(row, col)
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[row].height = 29
    guide.column_dimensions["B"].width = 76
    for col, width in (("B", 20), ("C", 19), ("D", 15), ("E", 23), ("F", 21), ("G", 13)):
        source.column_dimensions[col].width = width
    report.column_dimensions["B"].width = 27
    report.column_dimensions["C"].width = 18
    report.column_dimensions["D"].width = 27
    for row in range(4, 18):
        for col in (5, 6):
            source.cell(row, col).number_format = '#,##0'
    for row in range(4, 9):
        report.cell(row, 4).number_format = '#,##0'
    for ws, start, end in ((source, 4, 15), (report, 4, 6)):
        for row in range(start, end + 1):
            if row % 2 == 0:
                for cell in ws[row]:
                    cell.fill = PatternFill("solid", fgColor=pale)
    source.auto_filter.ref = "A3:G15"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output

def validate_saved_workbook(path: Path):
    wb = load_workbook(path)
    assert wb.sheetnames == ["Scope_And_Readme", "Synthetic_Inputs", "One_Report_Pilot"]
    ws, rep = wb["Synthetic_Inputs"], wb["One_Report_Pilot"]
    records = [tuple(ws.cell(i, j).value for j in range(1, 6)) for i in range(4, 16)]
    assert records == ROWS
    assert sum(r[3] * r[4] for r in records) == EXPECTED_TOMAN
    assert sum(r[3] for r in records) == EXPECTED_UNITS
    assert len({r[2] for r in records}) == len(records)
    assert all(ws.cell(i, 6).value == f"=D{i}*E{i}" for i in range(4, 16))
    assert rep["D4"].value.startswith("=SUMIF(")
    assert all(rep[f"B{i}"].value.startswith("=IF(") for i in (11,12,13))
    return {"sheets": len(wb.sheetnames), "records": len(records),
            "gross_toman_fictional": EXPECTED_TOMAN, "quantity_fictional": EXPECTED_UNITS,
            "formulas_verified": True, "excel_calculation_not_executed": True}

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("synthetic_one_report.xlsx"))
    args = parser.parse_args()
    path = build_workbook(args.output)
    result = validate_saved_workbook(path)
    print(f"PASS {path} | {result}")

if __name__ == "__main__":
    main()
